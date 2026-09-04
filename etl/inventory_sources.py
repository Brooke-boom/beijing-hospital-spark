# -*- coding: utf-8 -*-
"""子进程安全版数据文件盘点。每个文件在独立子进程解析，防单文件崩溃/OOM 拖垮全盘。
输出 /tmp/inventory_full.csv 与终端汇总。
"""
import os, sys, csv, glob, json, io, subprocess

ROOT = os.path.expanduser("~/Desktop/毕设/data")

INSPECT = r'''
# -*- coding: utf-8 -*-
import os, sys, csv, io, json, traceback
from openpyxl import load_workbook

def read_delimited(path):
    raw = open(path, 'rb').read()
    if raw[:2] == b'PK':
        return read_xlsx_bytes(raw)
    if raw[:4] == b'\xd0\xcf\x11\xe0':
        import xlrd
        wb = xlrd.open_workbook(file_contents=raw)
        sh = wb.sheet_by_index(0)
        rows = []
        for r in range(sh.nrows):
            rows.append(['' if c is None else str(c).strip() for c in sh.row_values(r)])
        wb.release_resources()
        return rows[0], rows[1:]
    for enc in ['utf-8-sig', 'utf-8', 'gb18030', 'utf-16']:
        try:
            text = raw.decode(enc); break
        except Exception: continue
    else:
        return None
    lines = [l for l in text.splitlines() if l.strip()]
    if not lines: return None
    def colcount(ch):
        try:
            return len(next(csv.reader([lines[0] + '\n' + (lines[1] if len(lines) > 1 else '')], delimiter=ch)))
        except Exception: return 1
    delim = max([',', ';', '\t'], key=colcount)
    try:
        rows = list(csv.reader(lines, delimiter=delim))
    except Exception:
        return None
    if not rows: return None
    return [h.strip().strip('\ufeff') for h in rows[0]], rows[1:]

def read_xlsx_bytes(raw):
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    try:
        header = [str(c).strip() if c is not None else '' for c in next(it)]
    except StopIteration:
        wb.close(); return None
    rows = []
    for row in it:
        if all(c is None or str(c).strip() == '' for c in row): continue
        rows.append(['' if c is None else str(c).strip() for c in row])
    wb.close()
    return header, rows

def load(path):
    ext = path.lower().endswith(('.xlsx', '.xlsm'))
    if ext or path.lower().endswith(('.xls', '.csv')):
        raw = open(path, 'rb').read(8)
    else:
        raw = b''
    if raw[:2] == b'PK':
        return read_xlsx_bytes(open(path, 'rb').read())
    if raw[:4] == b'\xd0\xcf\x11\xe0':
        import xlrd
        wb = xlrd.open_workbook(path)
        sh = wb.sheet_by_index(0)
        rows = []
        for r in range(sh.nrows):
            rows.append(['' if c is None else str(c).strip() for c in sh.row_values(r)])
        wb.release_resources()
        return rows[0], rows[1:]
    return read_delimited(path)

path = sys.argv[1]
try:
    res = load(path)
    if res is None:
        print(json.dumps({'ok': False, 'why': 'unreadable'}))
    else:
        h, rows = res
        print(json.dumps({'ok': True, 'rows': len(rows), 'cols': len(h), 'header': h}, ensure_ascii=False))
except Exception as e:
    print(json.dumps({'ok': False, 'why': repr(e)[:200]}, ensure_ascii=False))
'''

def main():
    out = []
    files = []
    for d in [ROOT, os.path.join(ROOT, '归档_旧版命名')]:
        files += sorted(glob.glob(os.path.join(d, '*.csv')) + glob.glob(os.path.join(d, '*.xlsx')) + glob.glob(os.path.join(d, '*.xls')))
    total = len(files)
    for i, f in enumerate(files):
        name = os.path.relpath(f, ROOT)
        try:
            cp = subprocess.run([sys.executable, '-c', INSPECT, f], capture_output=True, text=True, timeout=60, cwd=ROOT)
            line = cp.stdout.strip().splitlines()[-1] if cp.stdout.strip() else ''
            info = json.loads(line) if line.startswith('{') else {'ok': False, 'why': 'no-output'}
        except subprocess.TimeoutExpired:
            info = {'ok': False, 'why': 'timeout'}
        except Exception as e:
            info = {'ok': False, 'why': repr(e)[:150]}
        if not info.get('ok'):
            out.append({'file': name, 'rows': -1, 'cols': -1, 'grade': 0, 'bed': 0, 'phone': 0, 'web': 0, 'addr': 0, 'name': 0, 'fields': 'UNREADABLE: ' + info.get('why', '')})
            print(f"[{i+1}/{total}] ❌ {name:<46} {info.get('why','')[:60]}"); sys.stdout.flush()
            continue
        h = info['header']
        def has(*kws):
            return any(any(k in x for k in kws) for x in h)
        out.append({'file': name, 'rows': info['rows'], 'cols': info['cols'],
                    'grade': has('等级', '级别'), 'bed': has('床位'),
                    'phone': has('电话', '联系电话'), 'web': has('网站', '网址', '官网'),
                    'addr': has('地址'), 'name': has('名称', '机构'),
                    'fields': ' | '.join(x for x in h if x)[:260]})
        tags = []
        for t, k in [('等级', 'grade'), ('床位', 'bed'), ('电话', 'phone'), ('官网', 'web'), ('地址', 'addr')]:
            if out[-1][k]: tags.append(t)
        print(f"[{i+1}/{total}] {'✓' if tags else '-'} {name:<46} 行{info['rows']:>6} 列{info['cols']:>3}  {'、'.join(tags) if tags else '-'}")
        sys.stdout.flush()
    with open('/tmp/inventory_full.csv', 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.DictWriter(fh, fieldnames=['file', 'rows', 'cols', 'grade', 'bed', 'phone', 'web', 'addr', 'name', 'fields'])
        w.writeheader(); w.writerows(out)
    print('\n== 详情输出: /tmp/inventory_full.csv ==')

if __name__ == '__main__':
    main()
