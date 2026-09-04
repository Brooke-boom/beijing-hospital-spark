# -*- coding: utf-8 -*-
"""统一数据文件读取模块。自动识别: 真csv(utf-8/gb18030, ,/;/tab)、xlsx(含伪装.csv)、老.xls。
用法:
  import read_tables as rt; h, rows = rt.load(path)
  CLI: python read_tables.py <path>   # 打印 header + 前 N 行
"""
import os, sys, csv, io

def read_delimited(path):
    raw = open(path, 'rb').read()
    if raw[:2] == b'PK':
        return read_xlsx(path)
    if raw[:4] == b'\xd0\xcf\x11\xe0':
        return read_xls(path)
    for enc in ['utf-8-sig', 'utf-8', 'gb18030', 'utf-16']:
        try:
            text = raw.decode(enc)
            break
        except Exception:
            continue
    else:
        return None
    lines = [l for l in text.splitlines() if l.strip()]
    if not lines:
        return None

    def colcount(ch):
        try:
            sample = lines[0] + '\n' + (lines[1] if len(lines) > 1 else '')
            return len(next(csv.reader([sample], delimiter=ch)))
        except Exception:
            return 1
    delim = max([',', ';', '\t'], key=colcount)
    # 如果整个文件只有一列且以分号/特殊分隔，尝试重新按首行列数最多的分隔
    try:
        rows = list(csv.reader(lines, delimiter=delim))
    except Exception:
        return None
    if not rows:
        return None
    # 引用符包裹整行导致单列的情况: 再尝试 ',' 与 ';'
    if len(rows[0]) == 1 and len(lines) > 1 and (';' in lines[0] or ',' in lines[0]):
        for d2 in [';', ',']:
            if d2 == delim:
                continue
            try:
                rows2 = list(csv.reader(lines, delimiter=d2))
                if len(rows2[0]) > 1:
                    rows = rows2
                    break
            except Exception:
                continue
    return [h.strip().strip('\ufeff') for h in rows[0]], rows[1:]


def read_xlsx(path):
    from openpyxl import load_workbook
    raw = open(path, 'rb').read()
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    try:
        header = [str(c).strip() if c is not None else '' for c in next(it)]
    except StopIteration:
        wb.close()
        return None
    rows = []
    for row in it:
        if all(c is None or str(c).strip() == '' for c in row):
            continue
        rows.append(['' if c is None else str(c).strip() for c in row])
    wb.close()
    return header, rows


def read_xls(path):
    import xlrd
    wb = xlrd.open_workbook(path)
    sh = wb.sheet_by_index(0)
    rows = [['' if c is None else str(c).strip() for c in sh.row_values(r)] for r in range(sh.nrows)]
    wb.release_resources()
    return rows[0], rows[1:]


def load(path):
    low = path.lower()
    if low.endswith(('.xlsx', '.xlsm', '.xls')):
        return read_xls(path) if low.endswith('.xls') else read_xlsx(path)
    return read_delimited(path)


if __name__ == '__main__':
    path = sys.argv[1]
    n = int(sys.argv[2]) if len(sys.argv) > 2 else 8
    res = load(path)
    if res is None:
        print('UNREADABLE:', path)
        sys.exit(1)
    h, rows = res
    print('HEADER(%d): %s' % (len(h), ' | '.join(h)))
    print('ROWS: %d' % len(rows))
    for r in rows[:n]:
        print('  ', ' | '.join(r))
